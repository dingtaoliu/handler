"""One reader for every file the agent is handed: documents, spreadsheets, images.

Users send Excel, Word, PDF, PowerPoint, email files and photos (chat
attachments, Gmail attachments, Crontal downloads). Everything that turns such
a file into something a model can read goes through this module:

- `to_text(path)`: the whole file as Markdown, or None. Used where a tool saves
  a file and returns a text preview (crontal_api downloads, Gmail attachments,
  email intake).
- `read(path, ...)`: one bounded, pageable view for an agent's read_file tool.
  Text comes back at most `MAX_READ_CHARS` at a time with a note saying how to
  read the next part; images (and scanned or selected PDF pages) come back as
  image bytes the harness passes to the model as vision input.

Formats:
- Spreadsheets (.xlsx/.xlsm via openpyxl, .xls via xlrd, .csv/.tsv): every
  sheet, by name, as Markdown tables keyed by Excel column letters and row
  numbers, so a `sheet` + `cells` range ("A1:H200", "500:800") reads any part.
  Formulas show their last computed value.
- PDF, Word (.docx), PowerPoint (.pptx), HTML, EPUB, Outlook .msg, Jupyter,
  RSS, ZIP (members converted in turn), JSON, XML and plain text: markitdown.
  PDFs also take `pages`; a PDF without a text layer (a scan) comes back as page
  images, and `as_images` renders any pages (drawings, stamps) as images.
- Images (PNG, JPEG, GIF, WebP; HEIC, BMP, TIFF converted): downscaled to fit
  the model's vision limits.
- Anything else is decoded as UTF-8 if it is text, else reported as binary.

Legacy .doc/.ppt (pre-2007 binary Office) are not supported; the read says so.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import logging
import threading
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterator

from . import images as image_types

logger = logging.getLogger(__name__)

MAX_READ_CHARS = 20_000  # one read_file call's text (about 5k tokens)
MAX_TEXT_CHARS = 2_000_000  # to_text: whole-document conversions
MAX_CELL_CHARS = 500
MAX_PAGE_IMAGES = 4  # PDF pages rendered as images per call
PAGE_DPI = 110
IMAGE_MAX_SIDE = 2048  # longest side sent to the model
IMAGE_MAX_BYTES = 4 * 1024 * 1024
SCANNED_CHARS_PER_PAGE = 40  # fewer real characters per page: treat the PDF as a scan
SCAN_SAMPLE_PAGES = 12
# Printer registration marks that a scanned catalogue's text layer often holds.
_PRINTER_MARKS = {"C", "M", "Y", "K", "CM", "MY", "CY", "CMY", "CMYK"}
GARBAGE_RATIO = 0.05  # share of U+FFFD/control characters that makes text binary

SPREADSHEET_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".csv", ".tsv"}
# Converted by markitdown. Plain text and unknown suffixes are tried as UTF-8.
MARKITDOWN_SUFFIXES = {".pdf", ".docx", ".pptx", ".html", ".htm", ".xhtml", ".epub",
                       ".msg", ".ipynb", ".rss", ".atom", ".zip", ".json", ".jsonl", ".xml"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic", ".heif", ".bmp",
                  ".tif", ".tiff"}
UNSUPPORTED = {".doc": "legacy Word .doc", ".ppt": "legacy PowerPoint .ppt"}
DOCUMENT_SUFFIXES = SPREADSHEET_SUFFIXES | MARKITDOWN_SUFFIXES | {".txt", ".md"}


class ReadError(Exception):
    """A file the reader can't open; the message is meant for the agent."""


@dataclass
class Image:
    mime: str
    data: bytes
    label: str = ""


@dataclass
class Reading:
    text: str
    images: list[Image] = field(default_factory=list)


# ── entry points ───────────────────────────────────────────────────────────


def to_text(path: Path, *, max_chars: int = MAX_TEXT_CHARS) -> str | None:
    """The document as Markdown (up to about `max_chars`), or None.

    None means not a document, or conversion failed. Best effort, never raises:
    for text saved or previewed next to a downloaded file.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in DOCUMENT_SUFFIXES:
        return None
    try:
        if suffix in SPREADSHEET_SUFFIXES:
            if _spreadsheet_kind(path) == "html":
                return _cached(path, _markdown_html)[:max_chars]
            return _render_workbook(_open_workbook(path), budget=max_chars,
                                    name=path.name).text
        if suffix == ".pdf" and _looks_scanned(path):
            return None  # only printer marks or nothing: the agent should look at pages
        return _check_text(_cached(path, _markdown), path.name)[:max_chars]
    except Exception:  # noqa: BLE001 - conversion is best effort
        logger.info("could not convert %s", path.name, exc_info=True)
        return None


def read(path: Path, *, display: str | None = None, offset: int = 0,
         sheet: str | None = None, cells: str | None = None, pages: str | None = None,
         as_images: bool = False, max_chars: int = MAX_READ_CHARS) -> Reading:
    """A bounded view of `path` for the agent. Raises ReadError with a usable message.

    `display` is the path as the agent wrote it, used in continuation notes.
    """
    path = Path(path)
    name = display or path.name
    if not path.is_file():
        raise ReadError(f"not a file: {name}")
    suffix = path.suffix.lower()
    if suffix in UNSUPPORTED:
        raise ReadError(f"{name} is a {UNSUPPORTED[suffix]} file, which can't be read. "
                        "Ask for it as .docx/.pptx or PDF.")
    if image_types.sniff_file(path) or suffix in IMAGE_SUFFIXES:
        return _read_image(path, name)
    if suffix in SPREADSHEET_SUFFIXES:
        if _spreadsheet_kind(path) == "html":  # an HTML table saved as .xls (ERP exports)
            return _page(_convert(path, name, _markdown_html), offset, max_chars, name)
        return _render_workbook(_open_workbook(path), sheet=sheet, cells=cells,
                                budget=max_chars, name=name)
    if suffix == ".pdf":
        if as_images or _looks_scanned(path):
            # Never the whole file: at most MAX_PAGE_IMAGES pages per call.
            return _read_pdf_pages(path, name, pages, as_images=True, max_chars=max_chars,
                                   scanned=not as_images)
        if pages:
            return _read_pdf_pages(path, name, pages, as_images=False, max_chars=max_chars)
    if suffix in MARKITDOWN_SUFFIXES:
        return _page(_convert(path, name, _markdown), offset, max_chars, name)
    return _page(_plain_text(path, name), offset, max_chars, name)


def _convert(path: Path, name: str, convert: Callable[[Path], str]) -> str:
    try:
        text = _cached(path, convert)
    except Exception as exc:  # noqa: BLE001
        raise ReadError(f"could not convert {name}: {type(exc).__name__}: {exc}") from exc
    return _check_text(text, name)


# ── text paging ────────────────────────────────────────────────────────────


def _page(text: str, offset: int, max_chars: int, name: str) -> Reading:
    total = len(text)
    offset = max(0, offset)
    if offset >= total and total:
        return Reading(f"[{name}: offset {offset:,} is past the end ({total:,} characters).]")
    end = min(total, offset + max_chars)
    if end < total:
        # End on a line break when one is near, so a table row isn't split.
        cut = text.rfind("\n", offset + max_chars // 2, end)
        end = cut + 1 if cut > 0 else end
    chunk = text[offset:end]
    if offset or end < total:
        note = f"\n\n[{name}: characters {offset:,}–{end:,} of {total:,}."
        if end < total:
            note += f" Call read_file with offset={end} to read on."
        chunk += note + "]"
    return Reading(chunk or "(empty file)")


def _binary(name: str, size: int) -> ReadError:
    return ReadError(f"{name} is a binary file ({size:,} bytes) of a type that can't be "
                     "read as text; nothing was shown. If it is a document, ask for it as "
                     ".xlsx, .docx, .pptx or PDF.")


def _plain_text(path: Path, name: str) -> str:
    data = path.read_bytes()
    if b"\x00" in data[:8192]:
        raise _binary(name, len(data))
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if best is None:
            raise _binary(name, len(data)) from None
        text = str(best)
    return _check_text(text, name, len(data))


def _check_text(text: str, name: str, size: int | None = None) -> str:
    """Refuse text that is mostly decoding garbage (U+FFFD, NULs, control bytes)."""
    sample = text[:200_000]
    if not sample:
        return text
    bad = sum(1 for ch in sample
              if ch == "\ufffd" or (ord(ch) < 32 and ch not in "\t\n\r\f"))
    if bad / len(sample) > GARBAGE_RATIO:
        raise _binary(name, size if size is not None else len(text))
    return text


# ── markitdown ─────────────────────────────────────────────────────────────

_converter: Any = None
_converter_lock = threading.Lock()


def _markdown_converter() -> Any:
    global _converter
    with _converter_lock:
        if _converter is None:
            from markitdown import MarkItDown

            _converter = MarkItDown()
    return _converter


def _markdown(path: Path) -> str:
    return _markdown_converter().convert(str(path)).text_content or ""


def _markdown_html(path: Path) -> str:
    from markitdown import StreamInfo

    info = StreamInfo(extension=".html", mimetype="text/html")
    with path.open("rb") as f:
        return _markdown_converter().convert(f, stream_info=info).text_content or ""


_cache: OrderedDict[tuple, str] = OrderedDict()
_cache_lock = threading.Lock()
_CACHE_ENTRIES = 16


def _cached(path: Path, convert: Callable[[Path], str]) -> str:
    """Conversions kept per (file, size, mtime), so paging doesn't reconvert."""
    stat = path.stat()
    key = (str(path.resolve()), stat.st_size, stat.st_mtime_ns)
    with _cache_lock:
        if key in _cache:
            _cache.move_to_end(key)
            return _cache[key]
    text = convert(path)
    with _cache_lock:
        _cache[key] = text
        while len(_cache) > _CACHE_ENTRIES:
            _cache.popitem(last=False)
    return text


# ── spreadsheets ───────────────────────────────────────────────────────────


@dataclass
class _Sheet:
    name: str
    max_row: int | None
    max_col: int | None
    # rows(min_row, max_row, min_col, max_col) → (row number, values)
    rows: Callable[[int, int | None, int, int | None], Iterator[tuple[int, list[Any]]]]


@dataclass
class _Workbook:
    sheets: list[_Sheet]
    close: Callable[[], None] = lambda: None


def _spreadsheet_kind(path: Path) -> str:
    """From the bytes, since exports often lie: xlsx, xls, html, binary or text.

    Many ERP and bank exports named .xls are really HTML tables or CSV.
    """
    with path.open("rb") as f:
        head = f.read(2048)
    if head.startswith(b"PK\x03\x04"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    lower = head.lower()
    probe = lower.lstrip(b"\xef\xbb\xbf \t\r\n")
    if probe.startswith(b"<") and (b"<table" in lower or b"<html" in probe[:300]
                                   or b"<!doctype" in probe[:300]):
        return "html"
    if b"\x00" in head:
        return "binary"
    return "text"


def _open_workbook(path: Path) -> _Workbook:
    kind = _spreadsheet_kind(path)
    if kind == "xlsx":
        return _xlsx_workbook(path)
    if kind == "xls":
        return _xls_workbook(path)
    if kind == "binary":
        raise _binary(path.name, path.stat().st_size)
    return _csv_workbook(path, "\t" if path.suffix.lower() == ".tsv" else None)


def _xlsx_workbook(path: Path) -> _Workbook:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ReadError(f"could not open {path.name} as a workbook ({exc}); nothing was "
                        "shown. It may be corrupt or password-protected.") from exc
    sheets = []
    for ws in wb.worksheets:
        def rows(min_row, max_row, min_col, max_col, ws=ws):
            for i, row in enumerate(ws.iter_rows(min_row=min_row, max_row=max_row,
                                                 min_col=min_col, max_col=max_col,
                                                 values_only=True), start=min_row):
                yield i, list(row)
        sheets.append(_Sheet(ws.title, ws.max_row, ws.max_column, rows))
    return _Workbook(sheets, wb.close)


def _xls_workbook(path: Path) -> _Workbook:
    import xlrd

    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:  # noqa: BLE001
        raise ReadError(f"could not open {path.name} as an .xls workbook ({exc}); nothing "
                        "was shown. It may be corrupt or password-protected: ask for it "
                        "as .xlsx or PDF.") from exc
    sheets = []
    for index, title in enumerate(book.sheet_names()):
        sh = book.sheet_by_index(index)

        def rows(min_row, max_row, min_col, max_col, sh=sh):
            last_row = min(sh.nrows, max_row or sh.nrows)
            for r in range(min_row - 1, last_row):
                last_col = min(sh.ncols, max_col or sh.ncols)
                values = []
                for c in range(min_col - 1, last_col):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            values.append(xlrd.xldate.xldate_as_datetime(cell.value,
                                                                         book.datemode))
                            continue
                        except Exception:  # noqa: BLE001
                            pass
                    values.append(None if cell.ctype in (xlrd.XL_CELL_EMPTY,
                                                         xlrd.XL_CELL_BLANK) else cell.value)
                yield r + 1, values
        sheets.append(_Sheet(title, sh.nrows, sh.ncols, rows))
    return _Workbook(sheets, book.release_resources)


def _csv_workbook(path: Path, delimiter: str | None) -> _Workbook:
    text = _plain_text(path, path.name)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    width = max((len(r) for r in table), default=0)

    def rows(min_row, max_row, min_col, max_col):
        for r in range(min_row - 1, min(len(table), max_row or len(table))):
            yield r + 1, table[r][min_col - 1:max_col]
    return _Workbook([_Sheet(path.stem, len(table), width, rows)])


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() and abs(value) < 1e15 else repr(value)
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    text = str(value)
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS] + "…"
    return text.replace("\\", "\\\\").replace("|", "\\|").replace("\r", "").replace("\n", "<br>")


def _parse_cells(cells: str | None) -> tuple[int, int | None, int, int | None]:
    """'B2:F50', '100:200', 'C:E' or 'A10' → (min_row, max_row, min_col, max_col)."""
    if not cells:
        return 1, None, 1, None
    from openpyxl.utils import range_boundaries

    spec = cells.strip().upper().replace("$", "")
    if ":" not in spec:
        spec = f"{spec}:{spec}"
    try:
        min_col, min_row, max_col, max_row = range_boundaries(spec)
    except ValueError as exc:
        raise ReadError(f"bad cells range {cells!r}; use A1 notation like 'A1:H200', "
                        "'100:300' (rows) or 'C:F' (columns)") from exc
    return min_row or 1, max_row, min_col or 1, max_col


def _letter(col: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(col)


def _render_workbook(wb: _Workbook, *, sheet: str | None = None, cells: str | None = None,
                     budget: int | None, name: str = "") -> Reading:
    try:
        return _render(wb, sheet, cells, budget, name)
    finally:
        try:
            wb.close()
        except Exception:  # noqa: BLE001
            pass


def _dims(s: _Sheet) -> str:
    if s.max_row is None or s.max_col is None:
        return "size unknown"
    if not s.max_row or not s.max_col:
        return "empty"
    return f"A1:{_letter(s.max_col)}{s.max_row}"


def _render(wb: _Workbook, sheet: str | None, cells: str | None, budget: int | None,
            name: str) -> Reading:
    sheets = wb.sheets
    if sheet is not None:
        match = [s for s in sheets if s.name == sheet] or [
            s for s in sheets if s.name.strip().lower() == sheet.strip().lower()]
        if not match:
            raise ReadError(f"no sheet named {sheet!r}; sheets: "
                            + ", ".join(repr(s.name) for s in sheets))
        selected = match[:1]
    else:
        selected = sheets
    min_row, max_row, min_col, max_col = _parse_cells(cells)
    out: list[str] = []
    if len(sheets) > 1 or sheet is None:
        out.append(f"Workbook with {len(sheets)} sheet(s): "
                   + "; ".join(f"{s.name!r} ({_dims(s)})" for s in sheets))
    used = sum(len(x) + 1 for x in out)
    for index, s in enumerate(selected):
        heading = f"## Sheet: {s.name} ({_dims(s)})"
        rows: list[tuple[int, list[str]]] = []
        size = len(heading) + 2
        stopped_at: int | None = None
        for number, values in s.rows(min_row, max_row, min_col, max_col):
            texts = [_cell(v) for v in values]
            while texts and texts[-1] == "":
                texts.pop()
            if not texts:
                continue
            line = 8 + len(str(number)) + sum(len(t) + 3 for t in texts)
            if budget is not None and rows and used + size + line > budget:
                stopped_at = number
                break
            rows.append((number, texts))
            size += line
        out.append(heading)
        if rows:
            width = max(len(t) for _, t in rows)
            columns = [_letter(min_col + i) for i in range(width)]
            out.append("| row | " + " | ".join(columns) + " |")
            out.append("|---|" + "---|" * width)
            for number, texts in rows:
                out.append(f"| {number} | " + " | ".join(texts + [""] * (width - len(texts)))
                           + " |")
        else:
            out.append("(no values in this range)")
        used += size
        if stopped_at is not None:
            last_col = _letter(max_col) if max_col else (_letter(s.max_col) if s.max_col else "")
            rng = f"{_letter(min_col)}{stopped_at}:{last_col}{max_row or s.max_row or ''}"
            rest = [x.name for x in selected[index + 1:]]
            note = (f"[{name}: output limit reached in sheet {s.name!r} before row {stopped_at}"
                    f"{f' of {s.max_row}' if s.max_row else ''}. Call read_file with "
                    f"sheet={s.name!r} and cells={rng!r} to read on (a narrower column "
                    "range fits more rows)")
            if rest:
                note += "; not shown yet: " + ", ".join(repr(r) for r in rest)
            out.append(note + ".]")
            break
    return Reading("\n".join(out).rstrip() + "\n")


# ── PDF pages ──────────────────────────────────────────────────────────────


def _parse_pages(pages: str | None, count: int) -> list[int]:
    """'3', '2-5', '1,4,7-9' (1-based) → 0-based page indexes within the document."""
    if not pages:
        return list(range(count))
    out: list[int] = []
    try:
        for part in pages.replace(" ", "").split(","):
            if not part:
                continue
            if "-" in part:
                a, b = part.split("-", 1)
                start, end = int(a), int(b) if b else count
            else:
                start = end = int(part)
            out.extend(range(max(start, 1) - 1, min(end, count)))
    except ValueError as exc:
        raise ReadError(f"bad pages {pages!r}; use e.g. '3', '2-5' or '1,4,7-9'") from exc
    if not out:
        raise ReadError(f"pages {pages!r} are outside the document (it has {count} pages)")
    return list(dict.fromkeys(out))


def _looks_scanned(path: Path) -> bool:
    """Too little real text per page: a scan, or a catalogue whose text layer
    holds only printer marks ("C M Y CM MY CY CMY K"). Samples the first pages."""
    try:
        import fitz

        with fitz.open(path) as doc:
            sample = [doc[i].get_text("text")
                      for i in range(min(doc.page_count, SCAN_SAMPLE_PAGES))]
    except Exception:  # noqa: BLE001
        return False
    if not sample:
        return False
    real = sum(len(w) for text in sample for w in text.split()
               if w not in _PRINTER_MARKS and any(ch.isalnum() for ch in w))
    return real < SCANNED_CHARS_PER_PAGE * len(sample)


def _read_pdf_pages(path: Path, name: str, pages: str | None, *, as_images: bool,
                    max_chars: int, scanned: bool = False) -> Reading:
    import fitz

    try:
        doc = fitz.open(path)
    except Exception as exc:  # noqa: BLE001
        raise ReadError(f"could not open {name} as a PDF: {exc}") from exc
    with doc:
        wanted = _parse_pages(pages, doc.page_count)
        if as_images:
            shown = wanted[:MAX_PAGE_IMAGES]
            images = []
            for i in shown:
                pix = doc[i].get_pixmap(dpi=PAGE_DPI)
                images.append(_fit(Image("image/png", pix.tobytes("png"), f"page {i + 1}")))
            head = (f"{name} has no text layer (a scan); " if scanned else f"{name}: ")
            text = (head + f"pages {_ranges(shown)} of {doc.page_count} shown as images.")
            if len(wanted) > len(shown):
                rest = wanted[len(shown):]
                text += (f" Call read_file with pages='{_ranges(rest[:MAX_PAGE_IMAGES])}'"
                         f"{', as_images=true' if not scanned else ''} for the next ones.")
            return Reading(text, images)
        parts, used = [], 0
        for n, i in enumerate(wanted):
            body = doc[i].get_text("text").strip()
            block = f"--- page {i + 1} of {doc.page_count} ---\n{body}\n"
            if parts and used + len(block) > max_chars:
                rest = wanted[n:]
                parts.append(f"[{name}: output limit reached. Call read_file with "
                             f"pages='{_ranges(rest)}' to read on.]")
                break
            parts.append(block[:max_chars])
            used += len(block)
        return Reading("\n".join(parts))


def _ranges(indexes: list[int]) -> str:
    """0-based indexes → '1-3,7'."""
    out, start, prev = [], None, None
    for i in indexes:
        if start is None:
            start = prev = i
        elif i == prev + 1:
            prev = i
        else:
            out.append(f"{start + 1}" if start == prev else f"{start + 1}-{prev + 1}")
            start = prev = i
    if start is not None:
        out.append(f"{start + 1}" if start == prev else f"{start + 1}-{prev + 1}")
    return ",".join(out)


# ── images ─────────────────────────────────────────────────────────────────

_heif_registered = False


def _register_heif() -> None:
    global _heif_registered
    if not _heif_registered:
        _heif_registered = True
        try:
            from pillow_heif import register_heif_opener

            register_heif_opener()
        except Exception:  # noqa: BLE001 - optional; HEIC then reads as unsupported
            logger.info("pillow-heif unavailable; HEIC images can't be read")


def _read_image(path: Path, name: str) -> Reading:
    data = path.read_bytes()
    mime = image_types.sniff(data[:16])
    # A vision type within the limits goes as is; anything else is re-encoded.
    image = _fit(Image(mime or "", data, name), force=mime is None)
    return Reading(f"{name}: image ({_describe(data)}), shown below.", [image])


def _open_image(data: bytes) -> Any:
    from PIL import Image as PILImage

    _register_heif()
    return PILImage.open(io.BytesIO(data))


def _max_side(data: bytes) -> int:
    try:
        with _open_image(data) as img:
            return max(img.size)
    except Exception:  # noqa: BLE001
        return 0


def _describe(data: bytes) -> str:
    try:
        with _open_image(data) as img:
            return f"{img.format or 'image'}, {img.size[0]}×{img.size[1]}, {len(data):,} bytes"
    except Exception:  # noqa: BLE001
        return f"{len(data):,} bytes"


def _fit(image: Image, *, force: bool = False) -> Image:
    """Re-encode to PNG/JPEG within IMAGE_MAX_SIDE and IMAGE_MAX_BYTES if needed."""
    if (not force and len(image.data) <= IMAGE_MAX_BYTES
            and _max_side(image.data) <= IMAGE_MAX_SIDE):
        return image
    from PIL import ImageOps

    try:
        with _open_image(image.data) as img:
            img.seek(0)  # first frame of an animation
            frame = ImageOps.exif_transpose(img.copy())
    except Exception as exc:  # noqa: BLE001
        raise ReadError(f"{image.label or 'image'} could not be decoded as an image "
                        f"({type(exc).__name__})") from exc
    has_alpha = frame.mode in ("RGBA", "LA") or (
        frame.mode == "P" and "transparency" in frame.info)
    for side in (IMAGE_MAX_SIDE, IMAGE_MAX_SIDE // 2, IMAGE_MAX_SIDE // 4):
        if max(frame.size) > side:
            frame.thumbnail((side, side))
        buf = io.BytesIO()
        if has_alpha:
            frame.convert("RGBA").save(buf, "PNG", optimize=True)
            mime = "image/png"
        else:
            frame.convert("RGB").save(buf, "JPEG", quality=85)
            mime = "image/jpeg"
        if buf.tell() <= IMAGE_MAX_BYTES:
            return Image(mime, buf.getvalue(), image.label)
    raise ReadError(f"{image.label or 'image'} is too large to show even after downscaling")


# ── CLI (for the Codex harness's shell) ─────────────────────────────────────


def main(argv: list[str] | None = None) -> int:
    """handler-read FILE [--sheet S] [--cells R] [--pages P] [--offset N].

    Prints the same text the native harness's read_file returns. Page images
    (a scanned PDF, or --pages with --as-images) are saved under --images-dir
    and their paths printed, for Codex's view_image tool.
    """
    import argparse
    import sys

    parser = argparse.ArgumentParser(prog="handler-read", description=main.__doc__)
    parser.add_argument("path")
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--sheet")
    parser.add_argument("--cells")
    parser.add_argument("--pages")
    parser.add_argument("--as-images", action="store_true", help="render PDF pages as images")
    parser.add_argument("--images-dir", default=".tmp/handler-read")
    parser.add_argument("--max-chars", type=int, default=MAX_READ_CHARS)
    args = parser.parse_args(argv)
    source = Path(args.path)
    if image_types.sniff_file(source) or source.suffix.lower() in IMAGE_SUFFIXES:
        print(f"{args.path} is an image: open it with view_image.")
        return 0
    try:
        reading = read(source, display=args.path, offset=args.offset, sheet=args.sheet,
                       cells=args.cells, pages=args.pages, as_images=args.as_images,
                       max_chars=args.max_chars)
    except ReadError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    print(reading.text)
    if reading.images:
        out_dir = Path(args.images_dir)
        out_dir.mkdir(parents=True, exist_ok=True)
        for image in reading.images:
            ext = image_types.extension(image.mime) or ".png"
            target = out_dir / f"{source.stem}-{image.label.replace(' ', '')}{ext}"
            target.write_bytes(image.data)
            print(f"image: {target} (open it with view_image)")
    return 0


if __name__ == "__main__":  # pragma: no cover
    raise SystemExit(main())
