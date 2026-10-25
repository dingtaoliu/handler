"""The shared file reader (documents.py) and the native harness's read_file."""

from __future__ import annotations

import base64
import datetime as dt
import io
import json
import zipfile
from pathlib import Path

import fitz
import httpx
import openpyxl
import pytest
from agents import ItemHelpers, ToolOutputImage, ToolOutputText
from agents.tool_context import ToolContext
from PIL import Image

from handler.runtime import documents
from handler.runtime.documents import ReadError
from handler.runtime.gateway.crontal_api import _filename, _safe_name
from handler.runtime.harness import openai_agents as oa


# ── fixtures ───────────────────────────────────────────────────────────────


def _xlsx(path: Path, rows: int = 3) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quote"
    ws.append(["Item", "Grade", "Qty (t)", "Price | USD"])
    for i in range(rows):
        ws.append([f"flat bar {i}", "316L", 2.0 + i, 1234.5])
    ws["F1"] = dt.datetime(2026, 7, 31)
    other = wb.create_sheet("合同统计")
    other.append(["客户", "金额"])
    other.append(["Brava", 100])
    wb.save(path)
    return path


def _docx(path: Path, text: str) -> Path:
    files = {
        "[Content_Types].xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>"),
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>"),
        "word/document.xml": (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body></w:document>"),
    }
    with zipfile.ZipFile(path, "w") as z:
        for name, body in files.items():
            z.writestr(name, body)
    return path


def _pdf(path: Path, pages: list[str]) -> Path:
    doc = fitz.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    doc.save(path)
    doc.close()
    return path


def _png(path: Path, size=(40, 30), mode="RGB") -> Path:
    Image.new(mode, size, "red").save(path)
    return path


# ── spreadsheets ───────────────────────────────────────────────────────────


def test_xlsx_shows_every_sheet_as_tables(tmp_path):
    reading = documents.read(_xlsx(tmp_path / "q.xlsx"))
    text = reading.text
    assert not reading.images
    assert "Workbook with 2 sheet(s): 'Quote' (A1:F4); '合同统计' (A1:B2)" in text
    assert "## Sheet: Quote (A1:F4)" in text and "## Sheet: 合同统计 (A1:B2)" in text
    assert "| row | A | B | C | D | E | F |" in text
    assert "| 2 | flat bar 0 | 316L | 2 | 1234.5 |" in text
    assert "Price \\| USD" in text  # pipes in cells don't break the table
    assert "2026-07-31" in text
    assert "| 2 | Brava | 100 |" in text


def test_xlsx_sheet_and_cells_select_a_range(tmp_path):
    path = _xlsx(tmp_path / "q.xlsx", rows=20)
    text = documents.read(path, sheet="quote", cells="B10:C12").text
    assert "合同统计" in text.splitlines()[0]  # the sheet list is still shown
    assert "## Sheet: 合同统计" not in text
    assert "| row | B | C |" in text
    assert "| 10 | 316L | 10 |" in text and "| 12 | 316L | 12 |" in text
    assert "| 13 |" not in text and "| 9 |" not in text
    with pytest.raises(ReadError, match="no sheet named 'Nope'"):
        documents.read(path, sheet="Nope")
    with pytest.raises(ReadError, match="bad cells range"):
        documents.read(path, cells="not a range")


def test_large_sheet_is_cut_with_a_continuation(tmp_path):
    path = _xlsx(tmp_path / "big.xlsx", rows=500)
    text = documents.read(path, display="downloads/big.xlsx", max_chars=2000).text
    assert len(text) < 2600
    assert "output limit reached in sheet 'Quote' before row" in text
    assert "sheet='Quote' and cells='A" in text and ":F501'" in text
    assert "not shown yet: '合同统计'" in text
    row = int(text.split("before row ")[1].split(" ")[0])
    rest = documents.read(path, sheet="Quote", cells=f"A{row}:F501", max_chars=2000).text
    assert f"| {row} | flat bar {row - 2} |" in rest


def test_csv_with_semicolons(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("name;qty\nbar;3\npipe;4\n", encoding="utf-8")
    text = documents.read(path).text
    assert "| 1 | name | qty |" in text and "| 3 | pipe | 4 |" in text


def test_html_table_saved_as_xls_is_read(tmp_path):
    path = tmp_path / "export.xls"
    path.write_text("<html><body><table><tr><th>客户</th><th>金额</th></tr>"
                    "<tr><td>Brava</td><td>100</td></tr></table></body></html>",
                    encoding="utf-8")
    text = documents.read(path).text
    assert "客户" in text and "Brava" in text and "�" not in text


def test_undecodable_xls_is_refused_not_dumped(tmp_path):
    # A broken OLE file (the real failure dumped 396 KB of U+FFFD/NULs).
    path = tmp_path / "合同统计表-2026年7月31日.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(range(256)) * 400)
    with pytest.raises(ReadError) as info:
        documents.read(path)
    assert len(str(info.value)) < 400
    garbage = tmp_path / "blob.xlsm"
    garbage.write_bytes(b"\x01\x00\x02\x00" * 1000)
    with pytest.raises(ReadError, match="binary file"):
        documents.read(garbage)


# ── documents ──────────────────────────────────────────────────────────────


def test_docx_and_pptx(tmp_path):
    assert "316L flat bar quote" in documents.read(_docx(tmp_path / "a.docx",
                                                         "316L flat bar quote")).text
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Q3 supply plan"
    deck.save(tmp_path / "d.pptx")
    assert "Q3 supply plan" in documents.read(tmp_path / "d.pptx").text


def test_pdf_text_and_page_selection(tmp_path):
    words = " ".join(["stainless"] * 12)
    path = _pdf(tmp_path / "spec.pdf", [f"page one {words}", f"page two {words}",
                                        f"page three {words}"])
    full = documents.read(path)
    assert "page one" in full.text and "page three" in full.text and not full.images
    two = documents.read(path, pages="2-3").text
    assert "--- page 2 of 3 ---" in two and "page three" in two and "page one" not in two
    with pytest.raises(ReadError, match="outside the document"):
        documents.read(path, pages="9")


def test_scanned_pdf_comes_back_as_page_images(tmp_path):
    # Text layer holds only printer marks, like the 28.7 MB catalogue.
    path = _pdf(tmp_path / "catalogue.pdf", ["C M Y CM MY CY CMY K"] * 6)
    reading = documents.read(path, display="downloads/catalogue.pdf")
    assert len(reading.images) == documents.MAX_PAGE_IMAGES
    assert all(i.mime == "image/png" for i in reading.images)
    assert "no text layer" in reading.text and "pages 1-4 of 6" in reading.text
    assert "pages='5-6'" in reading.text
    rest = documents.read(path, pages="5-6")
    assert len(rest.images) == 2 and "pages 5-6 of 6" in rest.text
    assert documents.to_text(path) is None


def test_as_images_renders_text_pages(tmp_path):
    path = _pdf(tmp_path / "drawing.pdf", [" ".join(["dimension"] * 20)] * 2)
    reading = documents.read(path, pages="2", as_images=True)
    assert len(reading.images) == 1 and reading.images[0].label == "page 2"


# ── text and paging ────────────────────────────────────────────────────────


def test_long_text_pages_by_offset(tmp_path):
    path = tmp_path / "log.txt"
    content = "".join(f"line {i}\n" for i in range(10_000))
    path.write_text(content)
    first = documents.read(path, display="log.txt", max_chars=1000).text
    assert f"[log.txt: characters 0–" in first and f"of {len(content):,}." in first
    nxt = int(first.split("offset=")[1].split(" ")[0])
    last_line = first.split("\n\n")[0].splitlines()[-1]  # cut on a line break
    second = documents.read(path, offset=nxt, max_chars=1000).text
    assert second.splitlines()[0] == f"line {int(last_line.split()[1]) + 1}"
    assert "past the end" in documents.read(path, offset=10**7).text


def test_default_cap_is_about_20k(tmp_path):
    path = tmp_path / "big.md"
    path.write_text("x" * 100 + "\n" + "word " * 50_000)
    text = documents.read(path).text
    assert len(text) < documents.MAX_READ_CHARS + 300 and "offset=" in text


def test_binary_and_garbage_text_are_refused(tmp_path):
    blob = tmp_path / "blob.bin"
    blob.write_bytes(b"\x00\x01\x02" * 1000)
    with pytest.raises(ReadError, match="binary file"):
        documents.read(blob)
    noisy = tmp_path / "noisy.txt"
    noisy.write_text("ok " + "�" * 500)
    with pytest.raises(ReadError, match="binary file"):
        documents.read(noisy)


def test_gbk_text_is_decoded(tmp_path):
    path = tmp_path / "note.txt"
    path.write_bytes("不锈钢圆钢 316L 规格表，交货期三十天。".encode("gb18030") * 3)
    assert "不锈钢" in documents.read(path).text


def test_legacy_doc_is_explained(tmp_path):
    (tmp_path / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ReadError, match="legacy Word"):
        documents.read(tmp_path / "old.doc")


def test_to_text_previews(tmp_path):
    text = documents.to_text(_xlsx(tmp_path / "q.xlsx"))
    assert "## Sheet: 合同统计" in text
    assert documents.to_text(_png(tmp_path / "a.png")) is None
    assert documents.to_text(tmp_path / "missing.pdf") is None


# ── images ─────────────────────────────────────────────────────────────────


def test_small_image_goes_as_is(tmp_path):
    path = _png(tmp_path / "photo.png")
    reading = documents.read(path)
    assert reading.images[0].data == path.read_bytes()
    assert reading.images[0].mime == "image/png" and "40×30" in reading.text


def test_large_and_non_vision_images_are_converted(tmp_path):
    big = _png(tmp_path / "big.png", size=(5000, 1000))
    image = documents.read(big).images[0]
    with Image.open(io.BytesIO(image.data)) as img:
        assert max(img.size) == documents.IMAGE_MAX_SIDE
    bmp = tmp_path / "scan.bmp"
    Image.new("RGB", (20, 20)).save(bmp)
    assert documents.read(bmp).images[0].mime == "image/jpeg"
    named_wrong = tmp_path / "really_png.jpg"
    named_wrong.write_bytes(big.read_bytes()[:0] + _png(tmp_path / "s.png").read_bytes())
    assert documents.read(named_wrong).images[0].mime == "image/png"


def test_heic_is_converted_when_supported(tmp_path):
    pillow_heif = pytest.importorskip("pillow_heif")
    pillow_heif.register_heif_opener()
    path = tmp_path / "IMG_0001.HEIC"
    Image.new("RGB", (32, 32), "blue").save(path, format="HEIF")
    assert documents.read(path).images[0].mime == "image/jpeg"


# ── native harness read_file ───────────────────────────────────────────────


async def _call(ws: Path, **args):
    tool = next(t for t in oa.workspace_tools(ws) if t.name == "read_file")
    raw = json.dumps(args)
    ctx = ToolContext(context=None, tool_name="read_file", tool_call_id="c1",
                      tool_arguments=raw)
    return await tool.on_invoke_tool(ctx, raw)


async def test_read_file_returns_images_as_image_input(tmp_path):
    _png(tmp_path / "photo.png")
    out = await _call(tmp_path, path="photo.png")
    assert isinstance(out[0], ToolOutputText) and isinstance(out[1], ToolOutputImage)
    assert out[1].image_url.startswith("data:image/png;base64,")
    converted = ItemHelpers._convert_tool_output(out)
    assert [part["type"] for part in converted] == ["input_text", "input_image"]


async def test_read_file_reads_office_files_and_reports_errors(tmp_path):
    (tmp_path / "downloads").mkdir()
    _xlsx(tmp_path / "downloads" / "q.xlsx")
    out = await _call(tmp_path, path="downloads/q.xlsx", sheet="合同统计")
    assert isinstance(out, str) and "| 2 | Brava | 100 |" in out
    assert (await _call(tmp_path, path="../etc/passwd")).startswith("error: path escapes")
    assert (await _call(tmp_path, path="nope.pdf")).startswith("error: not a file")


def test_inline_image_data_is_kept_out_of_transcripts():
    url = "data:image/png;base64," + base64.b64encode(b"x" * 5000).decode()
    raw = {"type": "function_call_output", "call_id": "c1",
           "output": [{"type": "input_text", "text": "photo.png"},
                      {"type": "input_image", "image_url": url}]}
    cleaned = oa._raw(raw)
    assert cleaned["output"][0]["text"] == "photo.png"
    assert cleaned["output"][1]["image_url"].startswith("data:image/png;base64,…")
    assert len(json.dumps(cleaned)) < 400
    parts = oa._without_inline_data([ToolOutputImage(image_url=url)])
    assert "chars omitted" in parts[0]["image_url"]


# ── file names ─────────────────────────────────────────────────────────────


def test_names_keep_unicode():
    assert _safe_name("合同统计表-2026年7月31日.xls") == "合同统计表-2026年7月31日.xls"
    assert _safe_name("../../a/b\\c:d.pdf") == "c_d.pdf"
    assert _safe_name("bad\x00name\n.txt") == "bad_name_.txt"
    long = _safe_name("报" * 200 + ".xlsx")
    assert long.endswith(".xlsx") and len(long.encode()) <= 200
    assert _safe_name("..") == "download"


def test_download_name_prefers_rfc5987_filename():
    header = ("attachment; filename=-2026731.xls; filename*=UTF-8''"
              "%E5%90%88%E5%90%8C%E7%BB%9F%E8%AE%A1%E8%A1%A8-2026%E5%B9%B47%E6%9C%8831"
              "%E6%97%A5.xls")
    response = httpx.Response(200, headers={"content-disposition": header})
    assert _filename(response, "/api/v1/files/7") == "合同统计表-2026年7月31日.xls"
    plain = httpx.Response(200, headers={"content-disposition": 'attachment; filename="a b.pdf"'})
    assert _filename(plain, "/x") == "a b.pdf"


# ── CLI for Codex ──────────────────────────────────────────────────────────


def test_cli_prints_text_and_saves_page_images(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _xlsx(tmp_path / "q.xlsx")
    assert documents.main(["q.xlsx", "--sheet", "Quote", "--cells", "A1:B2"]) == 0
    assert "| 2 | flat bar 0 | 316L |" in capsys.readouterr().out
    _pdf(tmp_path / "scan.pdf", [""] * 2)
    assert documents.main(["scan.pdf"]) == 0
    out = capsys.readouterr().out
    assert "image: .tmp/handler-read/scan-page1.png" in out
    assert (tmp_path / ".tmp/handler-read/scan-page2.png").is_file()
    assert documents.main(["missing.xlsx"]) == 1
